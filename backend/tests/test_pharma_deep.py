"""Deep research trust boundaries, bounded rounds, recovery, attribution and real export formats."""
import asyncio,copy,io,json
from datetime import datetime,timezone
from unittest.mock import AsyncMock
import httpx,pytest
from pydantic import ValidationError
from src.services.pharma_deep_plan import DeepPlan,make_plan,search_calls,detect_type
from src.services.pharma_deep_engine import DeepResearch
from src.services.pharma_deep_evidence import add_records,coverage,verified_claim
from src.services.pharma_deep_report import tables_from,timeline_from,build_report
from src.services import pharma_deep_exports as exports,pharma_documents as documents
from src.services.llm_runtime import LLMProviderError

STAMP='2026-10-15T00:00:00+00:00'

def plan(**values):return DeepPlan(subjects=['sildenafil'],topics=['trials'],max_rounds=1,**values)
def record(**values):
    return {'title':'A registered sildenafil study','url':'https://clinicaltrials.gov/study/NCT01234567','kind':'trial_record',
        'facts':{'study_id':'NCT01234567','status':'COMPLETED','phases':['PHASE3'],'results_present':True,
            'completion':{'date':'2027-01-20','type':'ESTIMATED'},'primary_results':[{'measure':'Primary endpoint','time_frame':'Week 12','unit':'points','values':[{'group':'Treatment','value':'12.4','lower':'10.1','upper':'14.7'}]}]},
        'excerpt':'','date':'2026-10-14','date_meaning':'Registry update date','retrieved_at':STAMP,**values}
def job(p=None):return {'id':'aaaaaaaa-aaaa-4aaa-8aaa-aaaaaaaaaaaa','query':'Review the registered sildenafil study results and completion date.','plan':(p or plan()).model_dump(),'checkpoint':{},'activity':[],'started_at':datetime.now(timezone.utc)}

@pytest.mark.parametrize('question,kind',[
    ('Reblozyl drug profile','drug_profile'),('tirzepatide vs semaglutide head-to-head','head_to_head'),
    ('Pfizer obesity pipeline','company_pipeline'),('generic apixaban patents','generic_entry'),
    ('Keytruda safety warnings','safety'),('MDS competitor landscape','landscape')])
def test_question_specific_checklists(question,kind):assert detect_type(question)==kind

@pytest.mark.parametrize('values',[
    {'subjects':['https://evil.invalid/run']},{'subjects':['a\nsecret']},{'topics':['shell']},
    {'topics':['trials','trials']},{'max_rounds':4},{'question_type':'head_to_head'},
])
def test_edited_plans_cannot_introduce_unapproved_tools_or_unbounded_work(values):
    base={'subjects':['sildenafil'],'topics':['trials'],'max_rounds':3,**values}
    with pytest.raises(ValidationError):DeepPlan.model_validate(base)

@pytest.mark.asyncio
async def test_named_direct_comparisons_and_company_pipelines_do_not_spend_an_extra_planner_call(monkeypatch):
    from src.services import pharma_deep_plan as module
    model=AsyncMock();monkeypatch.setattr(module,'complete',model)
    direct=await make_plan('tirzepatide vs semaglutide head-to-head trials')
    calls=search_calls(direct)
    assert next(c for t,c in calls if t=='trials').compare_with=='semaglutide'
    pipeline=await make_plan('Pfizer obesity pipeline')
    assert next(c for t,c in search_calls(pipeline) if t=='trials').search_by=='sponsor'
    assert model.await_count==0

def test_results_numbers_are_verified_against_cited_quotes_not_other_sources():
    rows=[{**record(),'citation':1}]
    quotes=[{'id':'Qright','text':'The primary outcome value was 12.4 points.','source':1},
        {'id':'Qother','text':'Another study reported 98.7 points.','source':2}]
    good={'text':'The reported primary outcome value was 12.4 points.','quote_ids':['Qright']}
    assert verified_claim(good,quotes,rows,{'trials'})['sources']==[1]
    with pytest.raises((ValueError,LLMProviderError)):verified_claim({**good,'text':'The reported primary outcome value was 98.7 points.'},quotes,rows,{'trials'})
    with pytest.raises(ValueError):verified_claim({**good,'quote_ids':['fabricated']},quotes,rows,{'trials'})

@pytest.mark.parametrize('text',[
    'There are no trials for sildenafil.', 'Doctors may consider higher doses.',
    'Generic sildenafil can launch after the patent expires.', 'Study NCT99999999 completed in 2028-01-01.',
])
def test_unsupported_or_unsafe_claims_are_withheld(text):
    quotes=[{'id':'Qdata','text':'Sildenafil has a retrieved trial NCT01234567 with completion 2027-01-20 (ESTIMATED).','source':1}]
    with pytest.raises((ValueError,LLMProviderError)):
        verified_claim({'text':text,'quote_ids':['Qdata']},quotes,[{**record(),'citation':1}],{'trials'})

def test_source_rendered_tables_copy_results_and_preserve_estimated_milestones():
    tables=tables_from([{**record(),'citation':1}])
    study=next(t for t in tables if t['heading']=='Retrieved trials')
    outcome=next(t for t in tables if t['heading']=='Posted primary outcome values')
    assert '2027-01-20 (estimated)' in study['rows'][0]
    assert '12.4' in outcome['rows'][0] and 'Treatment' in outcome['rows'][0]

def test_deep_claim_preserves_estimated_date_and_does_not_invent_an_earliest_trial():
    rows=[{**record(),'citation':1}]
    quotes=[{'id':'Qdate','text':'NCT01234567: completion date 2027-01-20 (ESTIMATED).','source':1}]
    for text in ['NCT01234567 completed on 2027-01-20.','NCT01234567 has the earliest completion on 2027-01-20.']:
        with pytest.raises((ValueError,LLMProviderError)):verified_claim({'text':text,'quote_ids':['Qdate']},quotes,rows,{'trials'})

def test_current_indication_and_original_date_do_not_establish_the_original_population():
    quotes=[{'id':'Qold','text':'Initial authorisation: 2005-10-28.','source':1},{'id':'Qcurrent','text':'The current indication includes adults and children.','source':1}]
    with pytest.raises((ValueError,LLMProviderError)):
        verified_claim({'text':'The original authorisation covered adults and children.','quote_ids':['Qold','Qcurrent']},quotes,[],set())

def test_mds_phase_filters_are_broad_enough_for_pivotal_trials_and_nct_lookups_stay_exact():
    p=DeepPlan(subjects=['luspatercept','imetelstat'],topics=['trials'],condition='lower-risk MDS',phase='PHASE3')
    calls=search_calls(p)
    assert all(c.condition=='myelodysplastic syndromes' and c.phase=='PHASE3' for _,c in calls)
    exact=DeepPlan(subjects=['NCT01234567'],topics=['trials'],condition='obesity',phase='PHASE3')
    call=search_calls(exact)[0][1]
    assert call.query=='NCT01234567' and call.phase is None and call.condition is None

def test_todays_date_still_needs_a_supporting_quote():
    today=datetime.now(timezone.utc).date().isoformat()
    with pytest.raises((ValueError,LLMProviderError)):
        verified_claim({'text':'The regulatory action happened on '+today+'.','quote_ids':['Qstatus']},[{'id':'Qstatus','text':'The registry status is COMPLETED.','source':1}],[],set())

@pytest.mark.asyncio
async def test_workflow_publishes_only_verified_findings_and_resumes_without_repeating_search():
    search=AsyncMock(return_value={'status':'available','records':[record()],'coverage':'A bounded registry sample.'})
    writes=[]
    async def persist(state,activity,stage):writes.append(copy.deepcopy(state))
    async def model(messages,name,max_tokens):
        payload=json.loads(messages[-1]['content'])
        if isinstance(payload,list):return {'text':json.dumps({'approved':[row['id'] for row in payload]})}
        q=next(q for q in payload['quotes'] if q['field']=='status')
        return {'text':json.dumps({'claims':[{'text':'The retrieved study status is COMPLETED.','quote_ids':[q['id']]}]})}
    engine=DeepResearch(job(),persist,invoke=model,search=search,reader=AsyncMock())
    report,status=await engine.run()
    assert status=='completed' and report['metrics']['verified_claims']==1
    assert search.await_count==1 and report['tables'] and report['quotes']
    checkpoint=next(s for s in writes if s['phase']=='write')
    recovered=job();recovered['checkpoint']=checkpoint
    search.reset_mock()
    again=DeepResearch(recovered,persist,invoke=model,search=search,reader=AsyncMock())
    report,_=await again.run()
    assert search.await_count==0 and report['metrics']['model_calls']<=12
    recovered['checkpoint']=copy.deepcopy(again.state);recovered['checkpoint']['phase']='verify'
    model_forbidden=AsyncMock(side_effect=AssertionError('Verified work must not be repeated'))
    restored,_=await DeepResearch(recovered,persist,invoke=model_forbidden,search=search).run()
    assert restored['sections'] and model_forbidden.await_count==0

@pytest.mark.asyncio
async def test_failed_entailment_gets_one_repair_then_is_removed_not_published():
    search=AsyncMock(return_value={'status':'available','records':[record()],'coverage':'Registry sample.'})
    async def persist(*args):pass
    repairs=0
    async def model(messages,name,max_tokens):
        nonlocal repairs
        payload=json.loads(messages[-1]['content'])
        if isinstance(payload,list):return {'text':'{"approved":[]}'}
        if 'claims' in payload:repairs+=1;return {'text':'{"claims":[]}'}
        q=payload['quotes'][0]
        return {'text':json.dumps({'claims':[{'text':'The evidence definitively establishes commercial success.','quote_ids':[q['id']]}]})}
    engine=DeepResearch(job(),persist,invoke=model,search=search,reader=AsyncMock())
    report,status=await engine.run()
    assert repairs==1 and status=='partial' and not report['sections'] and report['not_verified']
    assert 'definitively establishes' not in report['markdown']

@pytest.mark.asyncio
async def test_budget_exhaustion_returns_source_evidence_without_unreviewed_model_text():
    search=AsyncMock(return_value={'status':'available','records':[record()],'coverage':'Registry sample.'})
    async def persist(*args):pass
    setup=job();setup['checkpoint']={'model_calls':12}
    model=AsyncMock()
    report,status=await DeepResearch(setup,persist,invoke=model,search=search,reader=AsyncMock()).run()
    assert status=='partial' and report['tables'] and not report['sections'] and model.await_count==0

@pytest.mark.asyncio
async def test_rounds_are_bounded_and_original_scope_is_not_overwritten_by_gap_queries():
    async def persist(*args):pass
    setup=job(DeepPlan(subjects=['sildenafil'],topics=['ema'],max_rounds=3))
    search=AsyncMock(return_value={'status':'empty','records':[],'coverage':'Empty retrieved sample, not proof of absence.'})
    model=AsyncMock(return_value={'text':'{"gaps":[{"topic":"ema","terms":"sildenafil assessment","reason":"Exact assessment not retrieved"}]}'} )
    engine=DeepResearch(setup,persist,invoke=model,search=search,reader=AsyncMock())
    report,status=await engine.run()
    assert report['metrics']['round']<=3 and search.await_count<=3 and status=='partial'
    assert engine.plan.subjects==['sildenafil']

@pytest.mark.asyncio
async def test_documents_have_real_exact_quotes_and_separate_provenance():
    async def persist(*args):pass
    text='Sildenafil is discussed in this official assessment report. The assessment describes the registered primary endpoint.'
    row=record(kind='indexed_regulatory',facts={},url='https://www.ema.europa.eu/en/documents/report.pdf',excerpt=text)
    search=AsyncMock(return_value={'status':'available','records':[row],'coverage':'Indexed excerpts only.'})
    reader=AsyncMock(return_value={'text':text,'url':row['url'],'read_method':'pdf_text','sha256':'a'*64,'retrieved_at':STAMP,'links':[],'truncated':False})
    setup=job(DeepPlan(subjects=['sildenafil'],topics=['ema'],max_rounds=1))
    engine=DeepResearch(setup,persist,search=search,reader=reader,invoke=AsyncMock(return_value={'text':'{"claims":[]}'}))
    report,_=await engine.run()
    passages=[q for q in report['quotes'] if q['method']=='document_passage']
    assert passages and all(q['text'] in text for q in passages) and report['sources'][0]['read_status']=='document_read'

@pytest.mark.parametrize('url',['http://www.fda.gov/a.pdf','https://fda.gov.evil.invalid/a','https://localhost/a','https://127.0.0.1/a','https://user:password@www.fda.gov/a','https://www.fda.gov:8443/a'])
def test_unapproved_document_locations_are_rejected(url):assert not documents.document_url(url)

@pytest.mark.asyncio
async def test_document_reader_revalidates_redirects_and_bounds_response_bytes(monkeypatch):
    async def resolve(url):
        if not documents.document_url(url):raise ValueError('Rejected host')
    transport=httpx.MockTransport(lambda request:httpx.Response(302,headers={'location':'https://127.0.0.1/private'}))
    async with httpx.AsyncClient(transport=transport) as client:
        with pytest.raises(ValueError):await documents.read_document('https://www.fda.gov/a',client=client,resolve=resolve)
    monkeypatch.setattr(documents,'MAX_BYTES',64)
    transport=httpx.MockTransport(lambda request:httpx.Response(200,headers={'content-type':'text/plain'},content=b'x'*65))
    async with httpx.AsyncClient(transport=transport) as client:
        with pytest.raises(ValueError):await documents.read_document('https://www.fda.gov/a',client=client,resolve=resolve)

def test_html_does_not_execute_or_quote_navigation_scripts_and_retains_only_exact_passages():
    parser=documents.ReadableHTML();parser.feed('<script>steal()</script><nav>menu</nav><p>Sildenafil trial results are posted in the public registry.</p>')
    text=parser.text();quotes=documents.relevant_quotes(text,'sildenafil trial')
    assert 'steal' not in text and 'menu' not in text and quotes==[text]

def sample_report():
    engine=DeepResearch(job(),AsyncMock());state=engine.state
    add_records(state,{'records':[record()],'status':'available'},'trials',job()['query'])
    state['coverage']=coverage(state,['trials'])
    return build_report('=HYPERLINK("https://evil.invalid")',plan(),state)

def test_word_pdf_and_excel_are_real_readable_exports_and_excel_never_executes_source_formulas():
    from docx import Document
    from openpyxl import load_workbook
    from pypdf import PdfReader
    report=sample_report();report['sources'][0]['title']='=HYPERLINK("https://evil.invalid")'
    doc=Document(io.BytesIO(exports.word(report)))
    assert any('Pharma deep research' in p.text for p in doc.paragraphs)
    pdf=PdfReader(io.BytesIO(exports.pdf(report)))
    assert len(pdf.pages)>=1 and 'Corten' in pdf.pages[0].extract_text()
    book=load_workbook(io.BytesIO(exports.excel(report)))
    assert book['Sources']['B2'].data_type=='s' and book['Sources']['B2'].value.startswith('=HYPERLINK')
    assert 'Evidence ledger' in book.sheetnames and 'Parsed API records' in book.sheetnames

@pytest.mark.asyncio
async def test_pdf_document_reader_extracts_real_text_in_a_separate_process():
    content=exports.pdf(sample_report())
    extracted=await documents.pdf_text(content)
    assert 'Primary endpoint' in extracted['text'] and extracted['pages_read']>=1 and not extracted['truncated']

def test_fda_timeline_sorts_calendar_dates_without_changing_their_source_values():
    rows=[{**record(),'citation':1,'date':'2026-11-02'},
        {**record(kind='fda_application',facts={'application':'NDA123456','approved_submissions':[{'date':'20261013','classification':'Labeling'}]}),'date':None,'citation':2}]
    timeline=timeline_from(rows)
    assert timeline[0]['date']=='20261013' and 'approved submission' in timeline[0]['meaning']
