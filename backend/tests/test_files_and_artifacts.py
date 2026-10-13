"""File creation (Word, PDF, Excel, CSV, PowerPoint) and artifacts."""
import io
import json
from unittest.mock import AsyncMock, patch
import pytest
from src.services import artifacts, file_builder
from src.services.chat_context import TurnContext
from src.services.llm_service import groq_llm_service
from src.services.search_orchestrator import SearchOrchestrator

EVIDENCE = [{"title": "EMA report", "url": "https://www.ema.europa.eu/r", "citationIndex": 1, "snippet": "14 approvals"},
            {"title": "Unused", "url": "https://example.org/u", "citationIndex": 2, "snippet": "x"}]
DOCUMENT = file_builder.DocumentSpec(title="Oncology / Briefing: 2026?", subtitle="For leadership", summary="A briefing.",
    sections=[{"heading": "Overview", "paragraphs": ["The EMA recommended 14 medicines [1]. 东京 Ünïcode"],
               "bullets": ["Lung cancer led"], "table": {"columns": ["Area", "Count"], "rows": [["Lung", 4], ["Breast", None]]}}])
WORKBOOK = file_builder.WorkbookSpec(title="Budget", sheets=[{"name": "Plan/2026", "columns": [
    {"header": "Month", "type": "date"}, {"header": "Spend", "type": "currency"}, {"header": "Share", "type": "percent"},
    {"header": "Note", "type": "text"}], "rows": [
    ["2026-01-01", 1200.5, 0.25, "ok"], ["2026-02-01", "1,300", "30%", "=HYPERLINK(\"http://evil\",\"x\")"],
    [None, "=SUM(B2:B3)", None, "=WEBSERVICE(\"http://x\")"]], "notes": "Sample values"}])
DECK = file_builder.DeckSpec(title="Pitch", subtitle="Pharma", slides=[
    {"title": "Problem", "bullets": ["Slow research [1]"], "notes": "Say this"},
    {"title": "Numbers", "table": {"columns": ["Metric", "Value"], "rows": [["Time saved", "40%"]]}}])


@pytest.mark.parametrize("query,kind", [
    ("Create a Word document summarizing Q3", "docx"),
    ("Can you make an Excel spreadsheet of our budget?", "xlsx"),
    ("Build a PowerPoint deck for the board", "pptx"),
    ("Generate slides about our roadmap", "pptx"),
    ("Export this as a PDF", "pdf"),
    ("Convert the table into a CSV", "csv"),
    ("Summarize this PDF for me", None),
    ("What is a PDF?", None),
    ("Write a report about EU pricing", None),
    ("How do I use Excel pivot tables?", None),
])
def test_detects_only_explicit_file_requests(query, kind):
    assert file_builder.detect(query) == kind


def test_documents_render_valid_files_with_cited_sources_only():
    from docx import Document
    from pypdf import PdfReader
    name, content_type, data = file_builder.render("docx", DOCUMENT, EVIDENCE)
    assert name == "Oncology Briefing 2026.docx" and content_type.endswith("wordprocessingml.document")
    text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    assert "[1] EMA report — https://www.ema.europa.eu/r" in text and "Unused" not in text
    name, content_type, data = file_builder.render("pdf", DOCUMENT, EVIDENCE)
    assert content_type == "application/pdf" and "EMA report" in PdfReader(io.BytesIO(data)).pages[0].extract_text()


def test_workbooks_keep_safe_formulas_and_neutralize_dangerous_ones():
    from openpyxl import load_workbook
    _, _, data = file_builder.render("xlsx", WORKBOOK, [])
    sheet = load_workbook(io.BytesIO(data))["Plan 2026"]
    assert sheet["B4"].value == "=SUM(B2:B3)" and sheet["B4"].data_type == "f"
    assert sheet["D3"].data_type == "s" and sheet["D4"].data_type == "s"   # HYPERLINK / WEBSERVICE stay text
    assert sheet["B3"].value == 1300 and sheet["C3"].value == 0.3 and sheet["C2"].number_format == "0.0%"
    assert sheet.freeze_panes == "A2" and sheet["A2"].value.year == 2026
    for formula, safe in (("=SUM(A1:A3)*2", True), ("=A1/B1", True), ("=HYPERLINK(\"x\")", False),
                          ("=WEBSERVICE(\"x\")", False), ("='[other.xlsx]S'!A1", False), ("=cmd|' /C calc'!A0", False)):
        assert file_builder.safe_formula(formula) is safe, formula


def test_csv_cells_can_never_execute_as_formulas():
    _, _, data = file_builder.render("csv", WORKBOOK, [])
    text = data.decode("utf-8-sig")
    assert "'=SUM(B2:B3)" in text and "'=HYPERLINK" in text and text.splitlines()[0] == "Month,Spend,Share,Note"


def test_presentations_include_cover_content_and_notes():
    from pptx import Presentation
    _, _, data = file_builder.render("pptx", DECK, EVIDENCE)
    deck = Presentation(io.BytesIO(data))
    assert len(deck.slides) == 4  # cover + 2 + sources
    assert deck.slides[1].notes_slide.notes_text_frame.text == "Say this"
    assert "40%" in [cell.text for shape in deck.slides[2].shapes if shape.has_table for row in shape.table.rows for cell in row.cells]
    assert "Pitch" in file_builder.describe("pptx", DECK) and "3 slides" in file_builder.describe("pptx", DECK)


def test_filenames_are_safe_on_every_operating_system():
    assert file_builder.filename('a/b\\c:d*e?"f<g>h|i', "docx") == "a b c d e f g h i.docx"
    assert file_builder.filename("   ...   ", "pdf") == "Corten file.pdf"
    assert len(file_builder.filename("x" * 300, "xlsx")) == 85


@pytest.mark.asyncio
async def test_invalid_model_output_gets_one_repair_attempt():
    class Completion:
        def __init__(self, content): self.choices = [type("C", (), {"message": type("M", (), {"content": content})()})()]
    replies = [Completion("not json"), Completion(json.dumps(DOCUMENT.model_dump()))]
    fake = type("Client", (), {})()
    fake.chat = type("Chat", (), {})()
    fake.chat.completions = type("Completions", (), {"create": AsyncMock(side_effect=replies)})()
    with patch.object(groq_llm_service, "client", fake):
        spec = await file_builder.generate_spec("docx", "Create a Word document", EVIDENCE, [])
    assert spec.title == DOCUMENT.title and fake.chat.completions.create.await_count == 2
    assert "SOURCES:" in fake.chat.completions.create.await_args_list[0].kwargs["messages"][0]["content"]


@pytest.mark.asyncio
async def test_file_requests_return_a_downloadable_file_event():
    route = {"strategy": "direct", "query_type": "creative", "tools": [], "requires_search": False}
    stored = {"id": "f1", "name": "Budget.xlsx", "kind": "generated", "size": 5000, "url": "/api/v1/attachments/f1/content"}
    with patch("src.services.search_orchestrator.smart_router.route", AsyncMock(return_value=route)), \
         patch("src.services.file_builder.generate_spec", AsyncMock(return_value=WORKBOOK)) as generate, \
         patch("src.services.attachments.store_generated", AsyncMock(return_value=stored)) as store, \
         patch("src.db.session.async_session_factory") as factory:
        factory.return_value.__aenter__.return_value.commit = AsyncMock()
        events = [json.loads(e[6:]) async for e in SearchOrchestrator().stream_search(
            "Make an Excel spreadsheet budget", turn=TurnContext(user_id="u1"))]
    file_event = next(e for e in events if e["type"] == "file")
    assert file_event["format"] == "xlsx" and file_event["url"] == stored["url"]
    assert store.await_args.args[2] == "Budget.xlsx" and generate.await_args.args[0] == "xlsx"
    assert "Budget" in "".join(e.get("content", "") for e in events if e["type"] == "token")
    assert events[-1]["type"] == "done"


def test_artifact_requests_and_streaming_guard():
    assert artifacts.requested("Build an interactive dashboard of monthly revenue")
    assert artifacts.requested("Draw a flowchart of our onboarding process")
    assert artifacts.requested("Create a React component for a pricing table")
    assert not artifacts.requested("What is the capital of France?")
    assert not artifacts.requested("Explain how dashboards help managers")
    opened = 'Intro\n:::artifact{type="html" title="X"}\n<h2>References</h2>\n'
    assert artifacts.inside_open_artifact(opened)
    assert not artifacts.inside_open_artifact(opened + "</html>\n:::\nSummary")
    assert not artifacts.inside_open_artifact("No artifact here")


@pytest.mark.asyncio
async def test_artifact_answers_use_pro_and_are_never_truncated():
    captured = {}

    async def answer(query, results, history, focus, **options):
        captured.update(options)
        yield 'Here it is.\n:::artifact{type="html" title="Page"}\n<h2>Sources</h2>\n'
        yield '<a href="https://a.example">A</a> <a href="https://b.example">B</a>\n'
        yield "</html>\n:::\nDone."

    route = {"strategy": "direct", "query_type": "code", "tools": [], "requires_search": False}
    with patch("src.services.search_orchestrator.smart_router.route", AsyncMock(return_value=route)), \
         patch.object(groq_llm_service, "stream_answer", answer):
        events = [json.loads(e[6:]) async for e in SearchOrchestrator().stream_search("Build a landing page")]
    text = "".join(e.get("content", "") for e in events if e["type"] == "token")
    assert text.endswith(":::\nDone.") and captured["artifact"] is True and captured["profile"].mode == "pro"
