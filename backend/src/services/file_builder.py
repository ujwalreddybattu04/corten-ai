"""Create real Word, PDF, Excel, CSV and PowerPoint files from a validated content spec.

The model writes structured JSON (never code); the server renders it. Facts come from the
supplied sources and keep their [n] citations, with a Sources section added automatically.
"""
import csv
import io
import json
import logging
import re
import unicodedata
from datetime import datetime
from typing import Literal, Optional, Union
from xml.sax.saxutils import escape
from pydantic import BaseModel, Field, ValidationError, field_validator
from src.config.settings import settings

logger = logging.getLogger(__name__)

FORMATS = {
    "docx": ("Word document", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "pdf": ("PDF document", "application/pdf"),
    "xlsx": ("Excel workbook", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "csv": ("CSV file", "text/csv"),
    "pptx": ("PowerPoint presentation", "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
}
ACTION = r"(?:create|generate|make|build|produce|prepare|draft|write|export|put|turn|convert|give me|save)"
PATTERNS = [
    ("pptx", r"\b(?:powerpoint|pptx|slide ?deck|slides|presentation|pitch deck)\b"),
    ("xlsx", r"\b(?:excel|xlsx|spreadsheet|workbook)\b"),
    ("csv", r"\bcsv\b"),
    ("docx", r"\b(?:word document|word doc|docx|\.doc\b|ms word|in word)\b"),
    ("pdf", r"\bpdf\b"),
]
ACCENT = "1F4E79"


def detect(query: str) -> str | None:
    """The file format the user explicitly asked to receive, or None for a normal answer."""
    text = query.lower()
    if not re.search(ACTION, text):
        return None
    for kind, pattern in PATTERNS:
        if re.search(pattern, text):
            # "Summarize this PDF" is about an input file, not a request for one.
            if kind == "pdf" and re.search(r"\b(?:this|the|attached|uploaded|my)\s+pdf\b", text) and not re.search(r"\b(?:as|into|to) (?:a )?pdf\b", text):
                return None
            return kind
    return None


# ── Content spec (bounded) ───────────────────────────────────────────────────

class TableSpec(BaseModel):
    columns: list[str] = Field(min_length=1, max_length=20)
    rows: list[list[Union[str, float, int, None]]] = Field(default_factory=list, max_length=200)


class SectionSpec(BaseModel):
    heading: str = Field(max_length=200)
    paragraphs: list[str] = Field(default_factory=list, max_length=30)
    bullets: list[str] = Field(default_factory=list, max_length=40)
    table: Optional[TableSpec] = None


class DocumentSpec(BaseModel):
    title: str = Field(min_length=1, max_length=200)
    subtitle: str = Field(default="", max_length=300)
    summary: str = Field(default="", max_length=1000)
    sections: list[SectionSpec] = Field(min_length=1, max_length=40)


class ColumnSpec(BaseModel):
    header: str = Field(min_length=1, max_length=100)
    type: Literal["text", "number", "currency", "percent", "date"] = "text"


class SheetSpec(BaseModel):
    name: str = Field(min_length=1, max_length=60)
    columns: list[ColumnSpec] = Field(min_length=1, max_length=40)
    rows: list[list[Union[str, float, int, None]]] = Field(default_factory=list, max_length=2000)
    notes: str = Field(default="", max_length=1000)


class WorkbookSpec(BaseModel):
    title: str = Field(min_length=1, max_length=200)
    summary: str = Field(default="", max_length=1000)
    sheets: list[SheetSpec] = Field(min_length=1, max_length=12)


class SlideSpec(BaseModel):
    title: str = Field(min_length=1, max_length=150)
    bullets: list[str] = Field(default_factory=list, max_length=8)
    notes: str = Field(default="", max_length=2000)
    table: Optional[TableSpec] = None

    @field_validator("bullets")
    @classmethod
    def short_bullets(cls, value):
        return [item[:300] for item in value]


class DeckSpec(BaseModel):
    title: str = Field(min_length=1, max_length=200)
    subtitle: str = Field(default="", max_length=300)
    summary: str = Field(default="", max_length=1000)
    slides: list[SlideSpec] = Field(min_length=1, max_length=40)


SPECS = {"docx": DocumentSpec, "pdf": DocumentSpec, "xlsx": WorkbookSpec, "csv": WorkbookSpec, "pptx": DeckSpec}
SCHEMAS = {
    "document": '{"title": str, "subtitle": str, "summary": "1-2 sentences for the chat", "sections": [{"heading": str, '
                '"paragraphs": [str], "bullets": [str], "table": {"columns": [str], "rows": [[str|number]]} | null}]}',
    "workbook": '{"title": str, "summary": "1-2 sentences for the chat", "sheets": [{"name": str, "columns": [{"header": str, '
                '"type": "text"|"number"|"currency"|"percent"|"date"}], "rows": [[value]], "notes": str}]}',
    "deck": '{"title": str, "subtitle": str, "summary": "1-2 sentences for the chat", "slides": [{"title": str, '
            '"bullets": [str, max 6, each under 120 chars], "notes": "speaker notes", "table": {...} | null}]}',
}


def spec_prompt(kind: str, has_sources: bool) -> str:
    family = {"docx": "document", "pdf": "document", "xlsx": "workbook", "csv": "workbook", "pptx": "deck"}[kind]
    label = FORMATS[kind][0]
    rules = [
        f"You write the content of a {label} that the user requested. Corten renders your JSON into the real file.",
        f"Return ONLY a JSON object with this shape: {SCHEMAS[family]}",
        "Write complete, specific, professional content an expert would deliver. Follow the user's requested structure, "
        "audience, language and length.",
        "Facts, figures and names must come from SOURCES when provided; cite them inline as [n]. Never invent statistics. "
        "If the user wants a template or there is no data, use clear placeholders or clearly labelled sample values.",
    ]
    if family == "workbook":
        rules.append("Numbers must be JSON numbers, not strings. Percent values are fractions (0.25 for 25%). Dates are "
                     "YYYY-MM-DD. Use Excel formulas (starting with '=') for totals and calculations, referencing cells in "
                     "the same sheet, e.g. \"=SUM(B2:B13)\". Row 1 is the header row, so data starts on row 2.")
    if family == "deck":
        rules.append("Aim for 6-12 slides: a clear storyline, one idea per slide, concise bullets, speaker notes with detail.")
    if not has_sources:
        rules.append("No sources were provided; rely on general knowledge only for well-established facts and avoid "
                     "precise current figures.")
    return "\n".join(rules)


def _sources_block(evidence: list[dict]) -> str:
    entries = []
    for index, source in enumerate(evidence[:12], start=1):
        number = source.get("citationIndex") or index
        entries.append(f"[{number}] {source.get('title', 'Source')} ({source.get('url', '')})\n{source.get('snippet', '')[:3000]}")
    return "SOURCES:\n" + "\n\n".join(entries) if entries else ""


async def generate_spec(kind: str, query: str, evidence: list[dict], history: list[dict], guidance: str = ""):
    """Ask the Pro model for a validated spec; one repair attempt on invalid output."""
    from src.services.llm_service import groq_llm_service
    from src.services.conversation_memory import answer_context
    system = spec_prompt(kind, bool(evidence)) + guidance + ("\n\n" + _sources_block(evidence) if evidence else "")
    messages = [{"role": "system", "content": system}]
    if history:
        messages.extend(answer_context(history, query))
    messages.append({"role": "user", "content": query[:16000]})
    last_error = ""
    for attempt in range(2):
        attempt_messages = messages + ([{"role": "user", "content": f"Your previous JSON was invalid ({last_error}). "
                                         "Return corrected JSON only."}] if attempt else [])
        completion = await groq_llm_service.client.chat.completions.create(
            messages=attempt_messages, model=settings.HEAVY_MODEL, temperature=0.3,
            max_tokens=settings.LLM_HEAVY_MAX_OUTPUT_TOKENS, reasoning_effort="medium",
            response_format={"type": "json_object"})
        try:
            return SPECS[kind].model_validate(json.loads(completion.choices[0].message.content or "{}"))
        except (json.JSONDecodeError, ValidationError) as error:
            last_error = type(error).__name__ + ": " + str(error)[:300]
            logger.info("File spec invalid on attempt %d (%s)", attempt + 1, type(error).__name__)
    return None


# ── Rendering ────────────────────────────────────────────────────────────────

def filename(title: str, kind: str) -> str:
    stem = unicodedata.normalize("NFKC", title)
    stem = re.sub(r'[\\/:*?"<>|\x00-\x1f]+', " ", stem)
    stem = re.sub(r"\s+", " ", stem).strip(" .")[:80] or "Corten file"
    return f"{stem}.{kind}"


def cited_sources(spec_text: str, evidence: list[dict]) -> list[tuple[int, dict]]:
    used = {int(n) for n in re.findall(r"\[(\d{1,2})\]", spec_text)}
    pairs = [((source.get("citationIndex") or index), source) for index, source in enumerate(evidence, start=1)]
    return [(number, source) for number, source in pairs if number in used and source.get("url", "").startswith("http")]


def render_docx(spec: DocumentSpec, sources: list[tuple[int, dict]]) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    for name in ("Title", "Heading 1", "Heading 2"):
        styles[name].font.color.rgb = RGBColor.from_string(ACCENT)
    document.add_paragraph(spec.title, style="Title")
    if spec.subtitle:
        document.add_paragraph(spec.subtitle, style="Subtitle")
    meta = document.add_paragraph(f"Prepared with Corten · {datetime.now().strftime('%d %B %Y')}")
    meta.runs[0].font.size, meta.runs[0].font.color.rgb = Pt(9), RGBColor(0x6B, 0x72, 0x80)
    for section in spec.sections:
        document.add_heading(section.heading, level=1)
        for paragraph in section.paragraphs:
            document.add_paragraph(paragraph)
        for bullet in section.bullets:
            document.add_paragraph(bullet, style="List Bullet")
        if section.table:
            _docx_table(document, section.table)
    if sources:
        document.add_heading("Sources", level=1)
        for number, source in sources:
            document.add_paragraph(f"[{number}] {source.get('title', 'Source')} — {source.get('url')}")
    # Page numbers in the footer.
    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), tag)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        run._r.append(element)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _docx_table(document, table: TableSpec):
    from docx.shared import RGBColor
    grid = document.add_table(rows=1, cols=len(table.columns))
    grid.style = "Light Grid Accent 1" if "Light Grid Accent 1" in [s.name for s in document.styles] else "Table Grid"
    for cell, header in zip(grid.rows[0].cells, table.columns):
        cell.text = str(header)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in table.rows:
        cells = grid.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = "" if value is None else str(value)
    document.add_paragraph()


def _pdf_markup(text: str) -> str:
    from reportlab.pdfbase import pdfmetrics
    from src.services.interpolation_pdf import registered_fonts
    fonts = registered_fonts()
    chunks, current, buffer = [], None, ""
    for char in unicodedata.normalize("NFC", str(text)):
        font = next((name for name in fonts if pdfmetrics.getFont(name).face.charToGlyph.get(ord(char), 0)), None)
        if font is None:
            char, font = "?", fonts[0]
        if current and current != font:
            chunks.append(f'<font name="{current}">{escape(buffer)}</font>')
            buffer = ""
        current, buffer = font, buffer + char
    if buffer:
        chunks.append(f'<font name="{current}">{escape(buffer)}</font>')
    return "".join(chunks)


def render_pdf(spec: DocumentSpec, sources: list[tuple[int, dict]]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from src.services.interpolation_pdf import registered_fonts
    font = registered_fonts()[0]
    accent = colors.HexColor("#" + ACCENT)
    body = ParagraphStyle("body", fontName=font, fontSize=10.5, leading=15, spaceAfter=6)
    title = ParagraphStyle("title", parent=body, fontSize=22, leading=27, textColor=accent, spaceAfter=4)
    subtitle = ParagraphStyle("subtitle", parent=body, fontSize=12.5, textColor=colors.HexColor("#4B5563"))
    heading = ParagraphStyle("heading", parent=body, fontSize=14, leading=18, textColor=accent, spaceBefore=12, spaceAfter=6)
    small = ParagraphStyle("small", parent=body, fontSize=8.5, textColor=colors.HexColor("#6B7280"))
    story = [Paragraph(_pdf_markup(spec.title), title)]
    if spec.subtitle:
        story.append(Paragraph(_pdf_markup(spec.subtitle), subtitle))
    story += [Paragraph(_pdf_markup(f"Prepared with Corten · {datetime.now().strftime('%d %B %Y')}"), small), Spacer(1, 8)]
    for section in spec.sections:
        story.append(Paragraph(_pdf_markup(section.heading), heading))
        story += [Paragraph(_pdf_markup(p), body) for p in section.paragraphs]
        if section.bullets:
            story.append(ListFlowable([ListItem(Paragraph(_pdf_markup(b), body), leftIndent=12) for b in section.bullets],
                                      bulletType="bullet", start="•", leftIndent=12))
        if section.table:
            data = [[Paragraph(_pdf_markup(c), small) for c in section.table.columns]] + [
                [Paragraph(_pdf_markup("" if v is None else v), small) for v in row[:len(section.table.columns)]]
                for row in section.table.rows]
            grid = Table(data, repeatRows=1, hAlign="LEFT")
            grid.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
                                      ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
                                      ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story += [grid, Spacer(1, 8)]
    if sources:
        story.append(Paragraph("Sources", heading))
        story += [Paragraph(_pdf_markup(f"[{n}] {s.get('title', 'Source')} — {s.get('url')}"), small) for n, s in sources]

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor("#6B7280"))
        canvas.drawCentredString(A4[0] / 2, 12 * mm, str(doc.page))
        canvas.restoreState()

    output = io.BytesIO()
    SimpleDocTemplate(output, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm, bottomMargin=20 * mm,
                      title=spec.title, author="Corten").build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


SAFE_FORMULA = re.compile(r"^=[A-Z0-9$:+\-*/^(),.%<>=&\" !]*$", re.I)
BLOCKED_FUNCTIONS = re.compile(r"\b(?:WEBSERVICE|FILTERXML|HYPERLINK|RTD|CALL|REGISTER|EXEC|DDE|INFO|CELL|IMPORT\w*)\s*\(", re.I)


def safe_formula(value: str) -> bool:
    """Formulas that only compute within the workbook; no links, external data or other workbooks."""
    return bool(SAFE_FORMULA.match(value)) and not BLOCKED_FUNCTIONS.search(value) and "[" not in value and "|" not in value


def _number(value, kind: str):
    if isinstance(value, (int, float)) or value is None:
        return value
    text = str(value).strip().replace(",", "").replace("€", "").replace("$", "").replace("£", "").replace("₹", "")
    try:
        if kind == "percent" and text.endswith("%"):
            return float(text[:-1]) / 100
        return float(text) if "." in text else int(text)
    except ValueError:
        return value


def render_xlsx(spec: WorkbookSpec, sources: list[tuple[int, dict]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    workbook = Workbook()
    workbook.remove(workbook.active)
    formats = {"number": "#,##0.##", "currency": "#,##0.00", "percent": "0.0%", "date": "yyyy-mm-dd"}
    used_names = set()
    for sheet_spec in spec.sheets:
        name = re.sub(r"[\[\]:*?/\\]", " ", sheet_spec.name)[:31].strip() or "Sheet"
        while name.lower() in used_names:
            name = (name[:28] + f" {len(used_names)}")[:31]
        used_names.add(name.lower())
        sheet = workbook.create_sheet(name)
        sheet.append([column.header for column in sheet_spec.columns])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=ACCENT)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row_index, row in enumerate(sheet_spec.rows, start=2):
            for column_index, column in enumerate(sheet_spec.columns, start=1):
                value = row[column_index - 1] if column_index - 1 < len(row) else None
                cell = sheet.cell(row=row_index, column=column_index)
                if isinstance(value, str) and value.startswith("="):
                    if safe_formula(value):
                        cell.value = value
                    else:
                        cell.value, cell.data_type = value, "s"  # shown as text, never evaluated
                elif column.type == "date" and isinstance(value, str):
                    try:
                        cell.value = datetime.strptime(value[:10], "%Y-%m-%d")
                    except ValueError:
                        cell.value = value
                elif column.type in ("number", "currency", "percent"):
                    cell.value = _number(value, column.type)
                else:
                    cell.value = value
                if column.type in formats:
                    cell.number_format = formats[column.type]
        for column_index, column in enumerate(sheet_spec.columns, start=1):
            longest = max([len(str(column.header))] + [len(str(r[column_index - 1])) for r in sheet_spec.rows
                                                         if column_index - 1 < len(r) and r[column_index - 1] is not None])
            sheet.column_dimensions[get_column_letter(column_index)].width = min(60, max(10, longest + 2))
        sheet.freeze_panes = "A2"
        if sheet_spec.rows:
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(sheet_spec.columns))}{len(sheet_spec.rows) + 1}"
        if sheet_spec.notes:
            note_row = len(sheet_spec.rows) + 3
            sheet.cell(row=note_row, column=1, value="Notes: " + sheet_spec.notes).font = Font(italic=True, color="6B7280")
    if sources:
        sheet = workbook.create_sheet("Sources")
        sheet.append(["#", "Source", "URL"])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for number, source in sources:
            sheet.append([number, source.get("title", "Source"), source.get("url")])
        sheet.column_dimensions["B"].width, sheet.column_dimensions["C"].width = 60, 80
    workbook.properties.title, workbook.properties.creator = spec.title, "Corten"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def render_csv(spec: WorkbookSpec, sources) -> bytes:
    sheet = spec.sheets[0]
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([column.header for column in sheet.columns])
    for row in sheet.rows:
        # Neutralize spreadsheet formula injection: CSV cells are always data.
        writer.writerow(["'" + v if isinstance(v, str) and v[:1] in "=+-@\t\r" else ("" if v is None else v)
                         for v in row[:len(sheet.columns)]])
    return ("﻿" + output.getvalue()).encode("utf-8")  # BOM so Excel opens Unicode correctly


def render_pptx(spec: DeckSpec, sources: list[tuple[int, dict]]) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    blank = deck.slide_layouts[6]
    navy, accent, ink, muted = RGBColor(0x0F, 0x17, 0x2A), RGBColor.from_string("3B82F6"), RGBColor(0x1F, 0x29, 0x37), RGBColor(0x6B, 0x72, 0x80)

    def text_box(slide, left, top, width, height, text, size, color, bold=False):
        frame = slide.shapes.add_textbox(left, top, width, height).text_frame
        frame.word_wrap = True
        frame.text = text
        run = frame.paragraphs[0].runs[0]
        run.font.size, run.font.bold, run.font.color.rgb = Pt(size), bold, color
        return frame

    cover = deck.slides.add_slide(blank)
    cover.background.fill.solid()
    cover.background.fill.fore_color.rgb = navy
    bar = cover.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(3.05), Inches(1.2), Inches(0.08))
    bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
    text_box(cover, Inches(0.8), Inches(1.6), Inches(11.5), Inches(1.4), spec.title, 40, RGBColor(0xFF, 0xFF, 0xFF), True)
    if spec.subtitle:
        text_box(cover, Inches(0.8), Inches(3.3), Inches(11.5), Inches(1), spec.subtitle, 20, RGBColor(0xCB, 0xD5, 0xE1))
    text_box(cover, Inches(0.8), Inches(6.5), Inches(8), Inches(0.5), datetime.now().strftime("%d %B %Y"), 12, RGBColor(0x94, 0xA3, 0xB8))

    for number, slide_spec in enumerate(spec.slides, start=2):
        slide = deck.slides.add_slide(blank)
        stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), deck.slide_height)
        stripe.fill.solid(); stripe.fill.fore_color.rgb = accent; stripe.line.fill.background()
        text_box(slide, Inches(0.6), Inches(0.4), Inches(12), Inches(1), slide_spec.title, 30, ink, True)
        top = Inches(1.55)
        if slide_spec.bullets:
            frame = slide.shapes.add_textbox(Inches(0.7), top, Inches(11.9 if not slide_spec.table else 5.6),
                                             Inches(5.2)).text_frame
            frame.word_wrap = True
            for index, bullet in enumerate(slide_spec.bullets):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.text = "•  " + bullet
                paragraph.space_after = Pt(10)
                paragraph.runs[0].font.size, paragraph.runs[0].font.color.rgb = Pt(20), ink
        if slide_spec.table:
            columns = slide_spec.table.columns
            rows = slide_spec.table.rows[:12]
            left = Inches(6.6) if slide_spec.bullets else Inches(0.7)
            width = Inches(6.1) if slide_spec.bullets else Inches(11.9)
            shape = slide.shapes.add_table(len(rows) + 1, len(columns), left, top, width, Inches(0.4) * (len(rows) + 1))
            for c, header in enumerate(columns):
                shape.table.cell(0, c).text = str(header)
            for r, row in enumerate(rows, start=1):
                for c in range(len(columns)):
                    value = row[c] if c < len(row) else ""
                    shape.table.cell(r, c).text = "" if value is None else str(value)
            for row in shape.table.rows:
                for cell in row.cells:
                    for paragraph in cell.text_frame.paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(12)
        text_box(slide, Inches(12.2), Inches(6.9), Inches(0.9), Inches(0.4), str(number), 11, muted)
        if slide_spec.notes:
            slide.notes_slide.notes_text_frame.text = slide_spec.notes
    if sources:
        slide = deck.slides.add_slide(blank)
        text_box(slide, Inches(0.6), Inches(0.4), Inches(12), Inches(1), "Sources", 30, ink, True)
        frame = slide.shapes.add_textbox(Inches(0.7), Inches(1.5), Inches(11.9), Inches(5.5)).text_frame
        frame.word_wrap = True
        for index, (number, source) in enumerate(sources[:12]):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = f"[{number}] {source.get('title', 'Source')} — {source.get('url')}"
            paragraph.runs[0].font.size, paragraph.runs[0].font.color.rgb = Pt(12), muted
    deck.core_properties.title, deck.core_properties.author = spec.title, "Corten"
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()


RENDERERS = {"docx": render_docx, "pdf": render_pdf, "xlsx": render_xlsx, "csv": render_csv, "pptx": render_pptx}


def render(kind: str, spec, evidence: list[dict]) -> tuple[str, str, bytes]:
    sources = cited_sources(spec.model_dump_json(), evidence)
    return filename(spec.title, kind), FORMATS[kind][1], RENDERERS[kind](spec, sources)


def describe(kind: str, spec) -> str:
    """A short chat message describing the created file."""
    if isinstance(spec, DeckSpec):
        outline = [slide.title for slide in spec.slides]
        detail = f"{len(spec.slides) + 1} slides"
    elif isinstance(spec, WorkbookSpec):
        outline = [f"{sheet.name} ({len(sheet.rows)} rows)" for sheet in spec.sheets]
        detail = f"{len(spec.sheets)} sheet{'s' if len(spec.sheets) != 1 else ''}"
    else:
        outline = [section.heading for section in spec.sections]
        detail = f"{len(spec.sections)} sections"
    lines = [spec.summary.strip() or f"Here is your {FORMATS[kind][0].lower()}.", "",
             f"**{spec.title}** · {FORMATS[kind][0]} · {detail}"]
    lines += [f"- {item}" for item in outline[:12]]
    return "\n".join(lines)
