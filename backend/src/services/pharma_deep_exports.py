"""Word/PDF reports and an Excel evidence workbook, generated from saved, verified data only."""
import io,json,re
from xml.sax.saxutils import escape

def plain(value):return re.sub(r'\*\*|(?<!\\)\*|\\([\\`*_{}\[\]<>#!|])',lambda m:m[1] if m.lastindex else '',str(value))

def blocks(report):
    yield 'title',report['title']
    yield 'paragraph',report['query']
    yield 'paragraph','Prepared '+report['generated_at']+' · Corten Pharma Intelligence'
    for section in report['sections']:
        yield 'heading',section['heading']
        for claim in section['claims']:yield 'paragraph',claim['text']+' '+''.join('['+str(s)+']' for s in claim['sources'])
    for table in report['tables']:yield 'heading',table['heading'];yield 'table',table
    if report['timeline']:
        yield 'heading','Timeline'
        yield 'table',{'columns':['Date','Record / milestone','Date meaning','Source'],'rows':[[t['date'],t['event'],t['meaning'],str(t['source'])] for t in report['timeline']]}
    yield 'heading','What remains unknown'
    for gap in report['unknown']:yield 'paragraph',gap['topic']+': '+gap['reason']
    if report['not_verified']:
        yield 'heading','Not verified'
        for gap in report['not_verified']:yield 'paragraph',gap['section']+': '+gap['reason']
    yield 'heading','Next checks'
    for check in report['next_checks']:yield 'paragraph',check
    yield 'paragraph',report['sample_notice']
    yield 'heading','Sources'
    for source in report['sources']:
        yield 'paragraph','['+str(source['id'])+'] '+source['title']+' · '+source['read_status']+'\n'+source['url']+'\nRetrieved '+source['retrieved_at']+(' · Source date '+str(source['date'])+' ('+str(source.get('date_meaning') or 'meaning not supplied')+')' if source.get('date') else '')

def word(report):
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    doc=Document();section=doc.sections[0];section.header.paragraphs[0].text='CORTEN  /  PHARMA INTELLIGENCE'
    section.top_margin=Inches(.8);section.bottom_margin=Inches(.8)
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10);normal.paragraph_format.space_after=Pt(9)
    for kind,data in blocks(report):
        if kind=='title':doc.add_heading(data,0)
        elif kind=='heading':doc.add_heading(data,1)
        elif kind=='paragraph':doc.add_paragraph(plain(data))
        else:
            table=doc.add_table(rows=1,cols=len(data['columns']));table.style='Light Shading Accent 1'
            for cell,text in zip(table.rows[0].cells,data['columns']):cell.text=plain(text)
            for row in data['rows']:
                for cell,text in zip(table.add_row().cells,row):cell.text=plain(text)
    section.footer.paragraphs[0].text='Source-backed research · Review original evidence before making decisions'
    output=io.BytesIO();doc.save(output);return output.getvalue()

def pdf(report):
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.lib.enums import TA_LEFT
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,LongTable,TableStyle
    from reportlab.lib.pagesizes import landscape,A4
    from src.services.interpolation_pdf import markup,registered_fonts
    registered_fonts()
    def unicode_text(value):return '<br/>'.join(markup(part) for part in plain(value).splitlines())
    output=io.BytesIO();page=landscape(A4);styles=getSampleStyleSheet()
    styles.add(ParagraphStyle('ResearchBody',fontName='Helvetica',fontSize=9,leading=14,spaceAfter=9,wordWrap='CJK'))
    styles.add(ParagraphStyle('ResearchCell',fontName='Helvetica',fontSize=7,leading=10,wordWrap='CJK'))
    story=[]
    for kind,data in blocks(report):
        if kind in ('title','heading','paragraph'):
            style=styles['Title'] if kind=='title' else styles['Heading2'] if kind=='heading' else styles['ResearchBody']
            story.append(Paragraph(unicode_text(data),style))
        else:
            grid=[[Paragraph(unicode_text(v),styles['ResearchCell']) for v in row] for row in [data['columns'],*data['rows']]]
            table=LongTable(grid,repeatRows=1,colWidths=[(page[0]-76)/len(data['columns'])]*len(data['columns']))
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e6edf1')),('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,0),.5,colors.HexColor('#8193a1')),('BOTTOMPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),7),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f5f7f8')])]))
            story.extend([table,Spacer(1,14)])
    def footer(canvas,doc):
        canvas.setFont('Helvetica',8);canvas.setFillColor(colors.HexColor('#61717e'));canvas.drawString(38,22,'Corten · Pharma Intelligence');canvas.drawRightString(page[0]-38,22,str(doc.page))
    SimpleDocTemplate(output,pagesize=page,rightMargin=38,leftMargin=38,topMargin=34,bottomMargin=38,title=report['title']).build(story,onFirstPage=footer,onLaterPages=footer)
    return output.getvalue()

def excel(report):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    book=Workbook();book.remove(book.active)
    def sheet(title,headers,rows):
        page=book.create_sheet(title[:31]);page.append(headers)
        for row in rows:
            page.append([str(v) if v is not None else '' for v in row])
        for row in page:
            for cell in row:
                # Force text, including leading '='/'+'/'-'/'@'; never create spreadsheet formulas from source text.
                cell.data_type='s';cell.alignment=Alignment(vertical='top',wrap_text=True)
        for cell in page[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='213544')
        page.freeze_panes='A2';page.auto_filter.ref=page.dimensions
        for column in page.columns:page.column_dimensions[column[0].column_letter].width=26 if column[0].column<4 else 48
    sheet('Findings',['Section','Verified finding','Quote IDs','Sources'],[[s['heading'],c['text'],', '.join(c['quote_ids']),', '.join(map(str,c['sources']))] for s in report['sections'] for c in s['claims']])
    sheet('Evidence ledger',['Quote ID','Source','Topic','Exact passage / field value','Method','Field','Document SHA256'],[[q['id'],q['source'],q['topic'],q['text'],q['method'],q.get('field'),(q.get('proof') or {}).get('sha256')] for q in report['quotes']])
    sheet('Sources',['Source','Title','URL','Source date','Date meaning','Retrieved UTC','Read status'],[[s['id'],s['title'],s['url'],s.get('date'),s.get('date_meaning'),s['retrieved_at'],s['read_status']] for s in report['sources']])
    sheet('Coverage',['Topic','Status','Sources','Quotes','Limitations'],[[c['label'],c['status'],c['sources'],c['quotes'],c['detail']] for c in report['coverage']])
    sheet('Not verified',['Section','Reason'],[[n['section'],n['reason']] for n in report['not_verified']])
    for index,table in enumerate(report['tables']):sheet(str(index+1)+' '+table['heading'],table['columns'],table['rows'])
    sheet('Timeline',['Date','Record / milestone','Date meaning','Source'],[[t['date'],t['event'],t['meaning'],t['source']] for t in report['timeline']])
    # Preserve all parsed records without silently losing long fields at Excel's cell limit.
    rows=[]
    for record in report['records']:
        data=json.dumps(record['facts'],ensure_ascii=False,sort_keys=True)
        for offset in range(0,len(data),30000):rows.append([record['citation'],record['kind'],record['title'],offset//30000+1,data[offset:offset+30000]])
    sheet('Parsed API records',['Source','Type','Title','Part','JSON source fields (continued by part)'],rows)
    output=io.BytesIO();book.save(output);return output.getvalue()
